"""Import the per-class data from the shared 'Personal Aion 2 Class Info' Google Sheet (Global season 1, by zxcastform)
into data/classes.json under classes.<id>.sheet, and add any skills it names that data/skills.json is missing.

Input: the sheet exported as xlsx (docs.google.com/spreadsheets/d/1lbHaVairHaz26M8XGNF9CBWqiiHHBM6CylTwNKnlOug/export?format=xlsx).
Usage: python tools/import_sheet_2026_10_02.py <book.xlsx> <allskills.json>
allskills.json = full aion2.app skill dump (id, en, ko, tw, k, q, icon, cd, mp ...) used to add missing skills.
"""
import json, os, re, sys, urllib.request

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK, ALL = sys.argv[1], sys.argv[2]
SHEETS = {"Chanter": "chanter", "Ranger": "ranger", "Cleric": "cleric", "Spirit Master": "spiritmaster", "Gladiator": "gladiator",
          "Sorcerer": "sorcerer", "Templar": "templar", "Assassin": "assassin"}
FIX = {  # sheet spelling -> game database name
    "Dark Rush": "Dark Crush", "Incandensent Blow": "Incandescent Blow", "Guardian's Blessing": "Guardian Blessing",
    "Impending Authority": "Impeding Authority", "Bitter Cold Wind": "Bittercold Wind", "Bitterwind Cold": "Bittercold Wind",
    "Winter Shackles": "Winter's Shackles", "Joinstrike: Corrode": "Jointstrike: Corrode", "Joinstrike Destruction": "Jointstrike: Destructive Attack",
    "Enhance: Spirit Benediction": "Enhance: Spirit's Benediction", "Siphone": "Siphon", "Elemental Unification": "Element Unification",
    "Claw of Beast": "Savage Fang", "Assult Stance": "Assault Stance", "ShadowFall": "Shadow Fall", "Focused Eyes": "Focused Eye",
    "Revilatization Contract": "Revitalization Contract", "Earth's Punishment": "Earth Punishment", "Experienced Counter": "Experienced Counterstrike",
    "Experienced Counterattack": "Experienced Counterstrike", "Hunters Resolve": "Hunter's Resolve", "Elemental Enhancement": "Element Enhancement",
    "Bitterwind": "Bittercold Wind", "Fire Spirit": "Summon: Fire Spirit", "Water Spirit": "Summon: Water Spirit", "Curse": "Jointstrike: Curse",
    "Judgement": "Judgment", "Smite": "Smite",
}
fix = lambda n: FIX.get(n.strip(), n.strip())

A = json.load(open(ALL, encoding="utf-8"))
SK = json.load(open(os.path.join(ROOT, "data/skills.json"), encoding="utf-8"))
DB = json.load(open(os.path.join(ROOT, "data/stigma_db.json"), encoding="utf-8"))
C = json.load(open(os.path.join(ROOT, "data/classes.json"), encoding="utf-8"))

def resolve(cls, name):
    """Return the skills.json key (adding the skill if needed) or the stigma name; None if unknown."""
    n = fix(name)
    if (cls + ":" + n) in DB:
        return n
    for key in (n, f"{n} ({cls})"):
        if key in SK and SK[key]["cls"] == cls:
            return key
    cand = [s for s in A[cls] if s["en"] == n and s.get("k") in ("Mastery", None) and s.get("d")]
    cand.sort(key=lambda s: (s.get("q") is None, -len(s.get("s") or [])))
    if not cand:
        return None
    s = cand[0]
    key = n if n not in SK else f"{n} ({cls})"
    icon = s.get("icon")
    if icon and not os.path.exists(os.path.join(ROOT, "icons", icon + ".webp")):
        try:
            img = urllib.request.urlopen(urllib.request.Request(f"https://aion2.app/db-item-icons/{icon}.webp", headers={"User-Agent": "Mozilla/5.0"}), timeout=30).read()
            open(os.path.join(ROOT, "icons", icon + ".webp"), "wb").write(img)
        except Exception:
            icon = None
    cd = s.get("cd")
    if isinstance(cd, float) and cd.is_integer():
        cd = int(cd)
    SK[key] = {"cls": cls, "id": s["id"], "tw": s["tw"], "ko": s["ko"], "cd": cd, "mp": s.get("mp"), "icon": icon}
    if "Passive" in (s.get("icon") or "") or (s.get("q") is None and not s.get("cd")):
        SK[key]["passive"] = 1
    print("  added skill", cls, key)
    return key

missing = []
wb = openpyxl.load_workbook(BOOK)
for title, cls in SHEETS.items():
    ws = wb[title]
    cell = lambda ref: (str(ws[ref].value).strip() if ws[ref].value is not None else "")
    rows = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                rows[c.coordinate] = str(c.value).strip()
    out = {"skills": [], "passives": [], "stigmas": [], "stigmaNotes": [], "cards": [], "gear": {}, "acc": {}, "genus": {}, "notes": []}
    # skill priority A3..A14: "Name (lvl) - (picks)"  or "Name - N/A"
    for r in range(3, 15):
        t = cell(f"A{r}")
        m = re.match(r"(.+?)\s*(?:\(([^)]*)\))?\s*-\s*\(?([\d/ ]+|N/A)\)?\s*$", t)
        if not m:
            continue
        name, lvl, picks = m.group(1), m.group(2) or "", m.group(3).strip()
        key = resolve(cls, name)
        if key is None:
            missing.append((cls, name))
        out["skills"].append([key or fix(name), lvl, [int(x) for x in re.findall(r"\d", picks)] if picks != "N/A" else []])
    # passives (B), stigmas (C)
    for r in range(3, 15):
        b, c = cell(f"B{r}"), cell(f"C{r}")
        mb = re.match(r"\d+\.\s*(.+)", b)
        if mb:
            key = resolve(cls, mb.group(1)) or fix(mb.group(1))
            if key == fix(mb.group(1)) and key not in SK and f"{key} ({cls})" not in SK:
                missing.append((cls, mb.group(1)))
            out["passives"].append(key)
        if c:
            mc = re.match(r"(\d+)\.\s*(.+)", c)
            if mc:
                key = resolve(cls, mc.group(2)) or fix(mc.group(2))
                if (cls + ":" + key) not in DB:
                    missing.append((cls, "stigma " + mc.group(2)))
                out["stigmas"].append([int(mc.group(1)), key])
            elif c.lower() == "or":
                continue
            elif len(c) > 40:
                out["stigmaNotes"].append(c)
            else:
                key = resolve(cls, c) or fix(c)
                if (cls + ":" + key) not in DB:
                    missing.append((cls, "stigma " + c))
                out["stigmas"].append([0, key])
    # Arcana / god stats / equip effects (D/E/F 3..12)
    for r in range(3, 13):
        d, e, f = cell(f"D{r}"), cell(f"E{r}"), cell(f"F{r}")
        if d:
            out["cards"].append([d, e, re.sub(r"^Target the follow(ing)?:\s*", "", f)])
    # labelled blocks anywhere on the sheet
    for ref, v in rows.items():
        col, r = re.match(r"([A-Z]+)(\d+)", ref).groups()
        nxt = rows.get(chr(ord(col) + 1) + r, "")
        if v in ("Weapon", "Guard", "Helmet", "Shoulders", "Chest", "Legs", "Gloves/Boots", "Cape") and col == "A":
            out["gear"][v] = nxt
        elif v in ("Earring & Necklace", "Rings", "Bracelets", "Brooches", "Seals") and col == "A":
            out["acc"][v] = nxt
        elif v == "Meta Equipped Wings:":
            out["wings"] = nxt
        elif v == "Highest Value Wings to level up:":
            out["wingsLevel"] = nxt
        elif v == "Enhancing Priority":
            out["enhance"] = nxt
        elif v == "Stats Priority":
            out["stones"] = nxt
        elif v in ("Offensive Titles", "Defensive Title", "Special Titles"):
            out.setdefault("titles", {})[v.replace(" Title", "").replace("s", "") if False else v] = nxt
        elif v == "Type" and col == "E":
            out["theostones"] = nxt
        elif v in ("Colossus", "Statues", "Paintings"):
            out.setdefault("pantheon", {})[v] = nxt
        elif col in "CD" and int(r) >= 50 and len(v) > 25 and not v.startswith("http"):
            out["notes"].append(v)
        elif col == "A" and int(r) >= 50 and len(v) > 25 and not v.startswith(("Credit", "http")):
            out["notes"].append(v)
        elif v.endswith("in-depth Guide:") and nxt.startswith("http"):
            out["guide"] = nxt
        elif v.endswith("guy (twitch/youtube):") and nxt.startswith("http"):
            out["creator"] = nxt
    # Genus boards
    for ref, v in rows.items():
        if v.endswith(" Board") and ref[0] in "ABC":
            col, r = ref[0], int(ref[1:])
            lines = []
            for k in range(r + 1, r + 7):
                t = rows.get(f"{col}{k}", "")
                if t.startswith("Line"):
                    lines.append(t)
            out["genus"][v] = lines
    out["notes"] = list(dict.fromkeys(out["notes"]))
    C["classes"][cls]["sheet"] = out
    print(cls, len(out["skills"]), "skills,", len(out["passives"]), "passives,", len(out["stigmas"]), "stigmas,", len(out["notes"]), "notes")

# shared sheets
sh = {"src": "https://docs.google.com/spreadsheets/d/1lbHaVairHaz26M8XGNF9CBWqiiHHBM6CylTwNKnlOug/htmlview",
      "author": "zxcastform (Discord)", "roadmap": [], "dailies": [], "buffs": []}
ws = wb["Roadmap"]
for r in range(1, ws.max_row + 1):
    a, b = ws[f"A{r}"].value, ws[f"B{r}"].value
    sh["roadmap"].append([str(a).strip() if a else "", str(b).strip() if b else ""])
ws = wb["DailyWeeklies"]
for r in range(1, ws.max_row + 1):
    vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else "" for c in range(1, 4)]
    if any(vals):
        sh["dailies"].append(vals)
ws = wb["Chanter+Cleric Buffs"]
for r in range(1, ws.max_row + 1):
    vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else "" for c in range(1, 9)]
    if any(vals):
        sh["buffs"].append(vals)
C["sheet"] = sh
json.dump(C, open(os.path.join(ROOT, "data/classes.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump(SK, open(os.path.join(ROOT, "data/skills.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)
print("unresolved:", missing)
